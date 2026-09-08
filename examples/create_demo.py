"""Cria ficheiros fictícios para experimentar a aplicação localmente."""

from pathlib import Path
from openpyxl import Workbook
from docx import Document


def create_demo(destination):
    root = Path(destination)
    templates = root / "modelos"
    templates.mkdir(parents=True, exist_ok=True)
    (root / "saida").mkdir(exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Doc Proposta"
    for cell, value in {
        "K3": "DEMO-001",
        "K4": "Reabilitação de edifício demonstrativo",
        "L6": "Local de exemplo",
        "N14": 90,
        "J25": "Documento demonstrativo",
        "J26": "Plano de trabalhos",
        "N45": "Fim do modelo demonstrativo",
    }.items():
        sheet[cell] = value
    workbook.save(root / "proposta-demo.xlsx")
    doc = Document()
    doc.add_heading("Proposta demonstrativa", 0)
    for text in [
        "Referência: {{ ID_INTERNO }}",
        "Empreitada: {{ NOME_DA_EMPREITADA }}",
        "Local: {{ LOCAL_DA_OBRA }}",
        "Prazo: {{ DURACAO_DA_EMPREITADA }} dias",
        "Documentos:\n{{ DOCUMENTOS_DA_PROPOSTA }}",
        "Data: {{ DATA_HOJE }}",
    ]:
        doc.add_paragraph(text)
    doc.add_paragraph("Exemplo fictício. Não constitui um modelo contratual.")
    doc.save(templates / "Proposta.docx")
    return root


if __name__ == "__main__":
    print(create_demo(Path(__file__).resolve().parents[1] / "local" / "demo"))
