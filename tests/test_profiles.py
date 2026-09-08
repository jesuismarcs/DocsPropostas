import copy
import json
from pathlib import Path
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName
from toolkit.profiles import default_profile, validate, portable, ProfileStore
from toolkit.mapping import Mapper, convert


def test_visual_mapping_reordered_headers_and_external_contacts(tmp_path):
    p = default_profile()
    p["items"].update(sheet="Medições", header_row=3, first_row=4)
    p["items"]["columns"] = {
        "ID": "Grupo",
        "Código": "Ref",
        "Descrição": "Artigo",
        "Quantidade": "Qtd",
        "Unidade": "Un",
    }
    p["items"]["types"] = {"Quantidade": "number"}
    p["contacts"]["columns"] = {"ID": "ID", "Email": "Emails"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Medições"
    ws.append(["Título"])
    ws.append([])
    ws.append(["Artigo", "Qtd", "Ref", "Un", "Grupo"])
    ws.append(["Pintura", "1.234,50", 7, "m2", "PINTURA"])
    ws["C4"].number_format = "0000"
    path = tmp_path / "mapa.xlsx"
    wb.save(path)
    contacts = Workbook()
    contacts.active.title = "Contactos"
    contacts.active.append(["ID", "Emails"])
    contacts.active.append(["PINTURA", "test@example.com"])
    source = tmp_path / "contactos.xlsx"
    contacts.save(source)
    p["sources"]["contacts"] = str(source)
    mapper = Mapper(p, path)
    try:
        assert mapper.table(p["items"]) == [
            {
                "ID": "PINTURA",
                "Código": "0007",
                "Descrição": "Pintura",
                "Quantidade": 1234.5,
                "Unidade": "m2",
            }
        ]
        assert mapper.table(p["contacts"])[0]["Email"] == "test@example.com"
        p["items"]["columns"]["Quantidade"] = "Quantidade ausente"
        with pytest.raises(ValueError, match="ausente"):
            mapper.table(p["items"])
    finally:
        mapper.close()


def test_named_ranges_tables_formulas_and_merged_cells(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(["Nome", "Valor"])
    ws.append(["Exemplo", 25])
    ws["D1"] = "=1+1"
    ws.merge_cells("F1:G1")
    ws["F1"] = "Unida"
    ws.add_table(Table(displayName="TabelaDados", ref="A1:B2"))
    wb.defined_names.add(DefinedName("Campo", attr_text="'Dados'!B2"))
    path = tmp_path / "dados.xlsx"
    wb.save(path)
    p = default_profile()
    mapper = Mapper(p, path)
    try:
        assert (
            mapper.field(
                {
                    "source": "main",
                    "sheet": "Dados",
                    "kind": "name",
                    "address": "Campo",
                    "type": "number",
                }
            )
            == 25
        )
        assert mapper.field(
            {
                "source": "main",
                "sheet": "Dados",
                "kind": "table",
                "address": "TabelaDados",
                "type": "table",
            }
        ) == [{"Nome": "Exemplo", "Valor": 25}]
        with pytest.raises(ValueError, match="recalcule"):
            mapper.field(
                {"sheet": "Dados", "kind": "cell", "address": "D1", "type": "number"}
            )
        p["rules"]["merged"] = "fill"
        assert (
            mapper.field(
                {"sheet": "Dados", "kind": "cell", "address": "G1", "type": "text"}
            )
            == "Unida"
        )
    finally:
        mapper.close()


def test_profile_export_sanitizes_without_changing_original(tmp_path):
    p = default_profile()
    p["sources"]["main"] = str(tmp_path / "private.xlsx")
    p["tools"]["email"]["smtp"].update(
        password="private", username="private@example.com"
    )
    p["proposal"]["CLIENTE"] = {
        "kind": "constant",
        "type": "text",
        "value": "Private client",
    }
    p["tools"]["excel"]["estaleiro"]["fields"][0]["value"] = "Private amount"
    serialized = json.dumps(portable(p))
    for secret in [
        "private.xlsx",
        "private@example.com",
        "Private client",
        "Private amount",
        "password",
    ]:
        assert secret not in serialized
    assert p["proposal"]["CLIENTE"]["value"] == "Private client"
    validate(portable(p))


def test_store_atomic_save_and_incompatible_import(tmp_path):
    store = ProfileStore("test", tmp_path)
    p = store.duplicate(store.active(), "Modelo alternativo")
    store.save(p)
    assert ProfileStore("test", tmp_path).active()["name"] == "Modelo alternativo"
    original = store.path.read_bytes()
    invalid = copy.deepcopy(p)
    invalid["version"] = 99
    with pytest.raises(ValueError):
        store.save(invalid)
    assert store.path.read_bytes() == original


def test_localized_numbers_dates():
    rules = default_profile()["rules"]
    assert convert("1.234,50", "number", rules) == 1234.5
    assert convert("31/12/2026", "date", rules) == "2026-12-31"
    assert (
        convert("1,234.50", "number", {**rules, "decimal": ".", "thousands": ","})
        == 1234.5
    )
    with pytest.raises(ValueError):
        convert("não é número", "number", rules)


def test_manual_lists_and_tables():
    mapper = Mapper(default_profile(), "")
    assert mapper.field({"kind": "manual", "type": "list", "value": "A\nB"}) == [
        "A",
        "B",
    ]
    assert mapper.field(
        {"kind": "manual", "type": "table", "value": "Nome\tValor\nExemplo\t10"}
    ) == [{"Nome": "Exemplo", "Valor": "10"}]
