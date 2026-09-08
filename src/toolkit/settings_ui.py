"""Shared native settings and workbook preview. Draft changes are explicit."""

import copy
import json
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from .profiles import validate, portable, default_profile
from .mapping import Workbook, Mapper


class SettingsWindow(tk.Toplevel):
    def __init__(self, parent, store, tool, on_saved=None):
        super().__init__(parent)
        self.title(
            "Definições · " + ("DocsPropostas" if tool == "docs" else "EnvioEmails")
        )
        self.geometry("1100x780")
        self.store, self.tool, self.on_saved = store, tool, on_saved
        self.draft = store.active()
        self.bindings = []
        self.book = None
        self.protocol("WM_DELETE_WINDOW", self.cancel)
        self.render()

    def cancel(self):
        if self.book:
            self.book.close()
        self.destroy()

    def render(self):
        for child in self.winfo_children():
            child.destroy()
        self.bindings = []
        bar = ttk.Frame(self, padding=10)
        bar.pack(fill="x")
        self.profile_name = tk.StringVar(value=self.draft["name"])
        ttk.Label(bar, text="Perfil").pack(side="left")
        selector = ttk.Combobox(
            bar,
            state="readonly",
            values=[p["name"] for p in self.store.data["profiles"]],
        )
        selector.set(self.draft["name"])
        selector.pack(side="left", padx=8)
        selector.bind("<<ComboboxSelected>>", lambda _: self.select(selector.current()))
        for text, action in [
            ("Novo", self.new),
            ("Duplicar", self.duplicate),
            ("Renomear", self.rename),
            ("Importar", self.import_profile),
            ("Exportar", self.export_profile),
            ("Restaurar padrão", self.restore),
        ]:
            ttk.Button(bar, text=text, command=action).pack(side="left", padx=2)
        self.tabs = ttk.Notebook(self)
        self.tabs.pack(fill="both", expand=True, padx=10)
        mapping = ttk.Frame(self.tabs, padding=8)
        self.tabs.add(mapping, text="Dados de entrada")
        self.mapping_page(mapping)
        rules = self.page("Regras e preferências")
        self.form(
            rules,
            self.draft["rules"],
            {
                "decimal": [",", "."],
                "thousands": [".", ",", "", " "],
                "hidden": ["include", "skip"],
                "merged": ["anchor", "fill", "error"],
                "blank_rows": ["skip", "stop"],
            },
        )
        self.form(rules, self.draft["preferences"], {"theme": ["light", "dark"]})
        output = self.page("Saídas")
        values = self.draft["tools"][self.tool]
        self.form(
            output,
            {k: v for k, v in values.items() if not isinstance(v, dict)},
            target=values,
        )
        if self.tool == "docs":
            catalog = self.page("Catálogo de habilitações")
            self.catalog_page(catalog)
        if self.tool == "email":
            for protocol in ("smtp", "imap"):
                frame = self.page(protocol.upper())
                self.form(
                    frame, values[protocol], {"security": ["ssl", "starttls", "plain"]}
                )
                var = tk.StringVar()
                ttk.Label(
                    frame, text="Nova palavra-passe (vazio mantém a guardada)"
                ).pack(anchor="w")
                ttk.Entry(frame, textvariable=var, show="•").pack(fill="x")
                setattr(self, protocol + "_password", var)
                ttk.Button(
                    frame,
                    text="Testar ligação sem enviar",
                    command=lambda p=protocol: self.test_connection(p),
                ).pack(anchor="w", pady=10)
        self.error = tk.StringVar()
        ttk.Label(
            self, textvariable=self.error, foreground="#a32020", wraplength=1000
        ).pack(fill="x", padx=10)
        actions = ttk.Frame(self, padding=10)
        actions.pack(fill="x")
        ttk.Button(actions, text="Cancelar", command=self.cancel).pack(side="right")
        ttk.Button(actions, text="Guardar e aplicar", command=self.save).pack(
            side="right", padx=8
        )

    def page(self, title):
        outer = ttk.Frame(self.tabs)
        self.tabs.add(outer, text=title)
        canvas = tk.Canvas(outer, highlightthickness=0)
        scroll = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        scroll.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        frame = ttk.Frame(canvas, padding=10)
        ident = canvas.create_window((0, 0), window=frame, anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        frame.bind(
            "<Configure>", lambda _: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(ident, width=e.width))
        return frame

    def form(self, parent, values, choices=None, target=None):
        target = target if target is not None else values
        labels = {
            "templates": "Pasta de modelos Word",
            "output": "Pasta de saída",
            "selected_models": "Modelos selecionados (; separa nomes; vazio = todos)",
            "filename": "Nome dos ficheiros (variáveis entre { })",
            "date_format": "Formato de data (%d/%m/%Y)",
            "alvara": "Utilizar habilitações",
            "decimal": "Separador decimal",
            "thousands": "Separador de milhares",
            "separator": "Separador de listas",
            "hidden": "Linhas ocultas",
            "merged": "Células unidas",
            "blank_rows": "Linhas vazias",
            "host": "Servidor",
            "port": "Porta",
            "security": "Segurança",
            "username": "Utilizador",
            "from": "Remetente",
            "auth": "Autenticar",
            "folder": "Pasta de correio",
            "base_folder": "Pasta base de obras",
            "deadline_days": "Antecedência do prazo (dias)",
            "subject": "Assunto",
            "signature": "Assinatura",
            "body": "Texto da mensagem",
            "reply_to": "Responder para",
            "recipient_mode": "Destinatários (to ou bcc)",
            "output_columns": "Colunas de saída (;)",
            "output_cell": "Primeira célula dos artigos",
            "output_sheet": "Folha de saída",
            "excel_template": "Modelo Excel (opcional)",
            "report_template": "Modelo Word do relatório (opcional)",
            "received_folder": "Subpasta das respostas",
            "language": "Idioma",
            "theme": "Tema",
        }
        for key, value in values.items():
            row = ttk.Frame(parent)
            row.pack(fill="x", pady=3)
            ttk.Label(row, text=labels.get(key, key), width=46).pack(side="left")
            var = (
                tk.BooleanVar(value=value)
                if isinstance(value, bool)
                else tk.StringVar(value=str(value))
            )
            if isinstance(value, bool):
                control = ttk.Checkbutton(row, variable=var)
            elif choices and key in choices:
                control = ttk.Combobox(
                    row, textvariable=var, values=choices[key], state="readonly"
                )
            else:
                control = ttk.Entry(row, textvariable=var)
            control.pack(side="left", fill="x", expand=True)
            if key in (
                "templates",
                "output",
                "base_folder",
                "excel_template",
                "report_template",
            ):
                ttk.Button(
                    row,
                    text="…",
                    width=3,
                    command=lambda v=var, k=key: self.choose_path(v, k),
                ).pack(side="right")
            self.bindings.append((target, key, var, type(value)))

    def choose_path(self, var, key):
        path = (
            filedialog.askopenfilename(parent=self)
            if key.endswith("_template")
            else filedialog.askdirectory(parent=self)
        )
        if path:
            var.set(path)

    def catalog_page(self, frame):
        from docspropostas.constants import ALVARA_DATA

        catalog = self.draft["tools"]["docs"]["catalog"]
        if not catalog:
            catalog.update(copy.deepcopy(ALVARA_DATA))
        table = ttk.Treeview(
            frame, columns=("category", "name", "class"), show="headings", height=16
        )
        for key, label in [
            ("category", "Categoria"),
            ("name", "Subcategoria"),
            ("class", "Classe"),
        ]:
            table.heading(key, text=label)
        table.pack(fill="both", expand=True)

        def refresh():
            table.delete(*table.get_children())
            for category, entries in catalog.items():
                for entry in entries:
                    table.insert(
                        "", "end", values=(category, entry["name"], entry["class"])
                    )

        variables = [tk.StringVar() for _ in range(3)]
        for label, var in zip(("Categoria", "Subcategoria", "Classe"), variables):
            ttk.Label(frame, text=label).pack(anchor="w")
            ttk.Entry(frame, textvariable=var).pack(fill="x")

        def selected(_):
            if table.selection():
                for var, value in zip(
                    variables, table.item(table.selection()[0], "values")
                ):
                    var.set(value)

        table.bind("<<TreeviewSelect>>", selected)

        def add():
            category, name, cls = [v.get().strip() for v in variables]
            if not category or not name:
                self.error.set("Indique categoria e subcategoria.")
                return
            entries = catalog.setdefault(category, [])
            entries[:] = [x for x in entries if x["name"] != name]
            entries.append({"name": name, "class": cls})
            refresh()

        def remove():
            category, name, _ = [v.get() for v in variables]
            if category in catalog:
                catalog[category][:] = [
                    x for x in catalog[category] if x["name"] != name
                ]
                if not catalog[category]:
                    del catalog[category]
            refresh()

        ttk.Button(frame, text="Adicionar / atualizar", command=add).pack(
            side="left", pady=10
        )
        ttk.Button(frame, text="Remover", command=remove).pack(side="left", padx=5)
        refresh()

    def collect(self):
        for obj, key, var, dtype in self.bindings:
            try:
                obj[key] = int(var.get()) if dtype is int else var.get()
            except ValueError as exc:
                raise ValueError(f"{key}: indique um número inteiro.") from exc
        return validate(self.draft)

    def select(self, index):
        self.draft = copy.deepcopy(self.store.data["profiles"][index])
        self.render()

    def new(self):
        self.draft = self.store.duplicate(default_profile(), "Novo perfil")
        self.render()

    def duplicate(self):
        self.draft = self.store.duplicate(
            self.collect(), self.draft["name"] + " — cópia"
        )
        self.render()

    def rename(self):
        name = simpledialog.askstring(
            "Nome", "Nome do perfil", initialvalue=self.draft["name"], parent=self
        )
        if name and name.strip():
            self.collect()
            self.draft["name"] = name.strip()
            self.render()

    def restore(self):
        self.draft = default_profile()
        self.render()

    def import_profile(self):
        path = filedialog.askopenfilename(
            parent=self, filetypes=[("Perfil JSON", "*.json")]
        )
        if path:
            try:
                with open(path, encoding="utf-8") as f:
                    candidate = validate(json.load(f))
                self.draft = self.store.duplicate(candidate, candidate["name"])
                self.render()
            except Exception as exc:
                self.error.set(str(exc))

    def export_profile(self):
        try:
            value = portable(self.collect())
            path = filedialog.asksaveasfilename(
                parent=self,
                defaultextension=".json",
                filetypes=[("Perfil JSON", "*.json")],
            )
            if path:
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(value, f, ensure_ascii=False, indent=2)
        except Exception as exc:
            self.error.set(str(exc))

    def save(self):
        try:
            self.collect()
            if self.tool == "email":
                for protocol in ("smtp", "imap"):
                    value = getattr(self, protocol + "_password").get()
                    if value:
                        self.store.set_password(self.draft["id"], protocol, value)
            self.store.save(self.draft)
            if self.on_saved:
                self.on_saved(copy.deepcopy(self.draft))
            self.cancel()
        except Exception as exc:
            self.error.set(str(exc))

    def mapping_page(self, frame):
        bar = ttk.Frame(frame)
        bar.pack(fill="x")
        ttk.Button(bar, text="Abrir Excel de exemplo…", command=self.open_book).pack(
            side="left"
        )
        self.sheet = ttk.Combobox(bar, state="readonly")
        self.sheet.pack(side="left", padx=8)
        self.sheet.bind("<<ComboboxSelected>>", lambda _: self.preview())
        ttk.Button(bar, text="Testar mapeamento", command=self.test_mapping).pack(
            side="right"
        )
        self.grid = ttk.Treeview(frame, show="headings", height=9)
        horizontal = ttk.Scrollbar(frame, orient="horizontal", command=self.grid.xview)
        vertical = ttk.Scrollbar(frame, orient="vertical", command=self.grid.yview)
        self.grid.configure(xscrollcommand=horizontal.set, yscrollcommand=vertical.set)
        vertical.pack(side="right", fill="y")
        self.grid.pack(fill="both", expand=True, pady=8)
        horizontal.pack(fill="x")
        self.grid.bind("<ButtonRelease-1>", self.pick_cell)
        bottom = ttk.Frame(frame)
        bottom.pack(fill="x")
        self.section = tk.StringVar(value="proposal")
        section = ttk.Combobox(
            bottom,
            textvariable=self.section,
            values=["proposal", "items", "contacts"],
            state="readonly",
            width=14,
        )
        section.pack(side="left")
        section.bind("<<ComboboxSelected>>", lambda _: self.refresh_fields())
        self.field_name = ttk.Combobox(bottom, width=32)
        self.field_name.pack(side="left", padx=4)
        self.field_name.bind("<<ComboboxSelected>>", lambda _: self.load_mapping())
        ttk.Button(bottom, text="Adicionar campo", command=self.add_field).pack(
            side="left"
        )
        ttk.Button(bottom, text="Remover campo", command=self.remove_field).pack(
            side="left"
        )
        self.editor = ttk.Frame(frame)
        self.editor.pack(fill="x", pady=8)
        self.refresh_fields()

    def refresh_fields(self):
        section = self.section.get()
        self.field_name["values"] = (
            list(self.draft[section])
            if section == "proposal"
            else list(self.draft[section]["columns"])
        )
        if self.field_name["values"]:
            self.field_name.current(0)
        self.load_mapping()

    def load_mapping(self):
        for child in self.editor.winfo_children():
            child.destroy()
        section, key = self.section.get(), self.field_name.get()
        if section == "proposal" and key not in self.draft[section]:
            return
        m = self.draft[section][key] if section == "proposal" else self.draft[section]
        self.map_vars = {}
        fields = [
            ("source", "Ficheiro", m.get("source", "main")),
            ("sheet", "Folha", m.get("sheet", "")),
            ("kind", "Origem", m.get("kind", "range")),
            ("address", "Célula / intervalo / nome", m.get("address", "")),
        ]
        if section == "proposal":
            fields += [
                ("type", "Tipo", m.get("type", "text")),
                ("value", "Valor constante/manual", m.get("value", "")),
                (
                    "required",
                    "Obrigatório (sim/não)",
                    "sim" if m.get("required") else "não",
                ),
            ]
        else:
            fields += [
                ("header_row", "Linha cabeçalhos", m.get("header_row", 1)),
                ("first_row", "Primeira linha", m.get("first_row", 2)),
                ("last_row", "Última linha (0 = fim)", m.get("last_row", 0)),
                ("column", "Cabeçalho ou @coluna", m["columns"].get(key, "")),
                ("column_type", "Tipo da coluna", m.get("types", {}).get(key, "text")),
            ]
        for i, (name, label, value) in enumerate(fields):
            ttk.Label(self.editor, text=label).grid(
                row=i // 2, column=(i % 2) * 2, sticky="w"
            )
            var = tk.StringVar(value=str(value))
            self.map_vars[name] = var
            choices = {
                "kind": ["cell", "range", "name", "table", "constant", "manual"]
                if section == "proposal"
                else ["range", "name", "table"],
                "type": ["text", "number", "date", "list", "table"],
                "source": list(self.draft["sources"]),
            }
            choices.update(
                required=["sim", "não"], column_type=["text", "number", "date"]
            )
            widget = (
                ttk.Combobox(
                    self.editor,
                    textvariable=var,
                    values=choices[name],
                    state="readonly",
                )
                if name in choices
                else ttk.Entry(self.editor, textvariable=var)
            )
            widget.grid(row=i // 2, column=(i % 2) * 2 + 1, sticky="ew", padx=5, pady=2)
        row = (len(fields) + 1) // 2
        ttk.Button(
            self.editor, text="Aplicar associação", command=self.apply_mapping
        ).grid(row=row, column=0, pady=8)
        ttk.Button(
            self.editor, text="Associar ficheiro de origem…", command=self.bind_source
        ).grid(row=row, column=1)
        ttk.Button(self.editor, text="Nova origem", command=self.add_source).grid(
            row=row, column=2
        )
        if section == "proposal":
            ttk.Button(
                self.editor, text="Editar valor longo…", command=self.edit_value
            ).grid(row=row, column=3)
        self.editor.columnconfigure(1, weight=1)
        self.editor.columnconfigure(3, weight=1)

    def apply_mapping(self):
        try:
            section, key = self.section.get(), self.field_name.get()
            m = (
                self.draft[section][key]
                if section == "proposal"
                else self.draft[section]
            )
            for name, var in self.map_vars.items():
                if name == "column":
                    m["columns"][key] = var.get()
                elif name == "column_type":
                    m.setdefault("types", {})[key] = var.get()
                elif name == "required":
                    m[name] = var.get() == "sim"
                else:
                    m[name] = (
                        int(var.get())
                        if name in ("header_row", "first_row", "last_row")
                        else var.get()
                    )
            self.error.set("Associação aplicada ao rascunho. Guarde para utilizar.")
        except Exception as exc:
            self.error.set(str(exc))

    def add_source(self):
        name = simpledialog.askstring("Origem", "Nome da origem de dados", parent=self)
        if name and name.strip() and name not in self.draft["sources"]:
            self.draft["sources"][name] = ""
            self.load_mapping()

    def edit_value(self):
        dialog = tk.Toplevel(self)
        dialog.title(
            "Valor · listas: uma linha por item · tabelas: colunas separadas por tabulação"
        )
        text = tk.Text(dialog, width=90, height=18)
        text.pack(fill="both", expand=True)
        text.insert("1.0", self.map_vars["value"].get())

        def save():
            self.map_vars["value"].set(text.get("1.0", "end").rstrip())
            dialog.destroy()

        ttk.Button(dialog, text="Aplicar valor", command=save).pack(pady=8)

    def bind_source(self):
        path = filedialog.askopenfilename(
            parent=self, filetypes=[("Excel", "*.xlsx *.xlsm *.xls")]
        )
        if path:
            self.draft["sources"][self.map_vars["source"].get()] = path

    def add_field(self):
        name = simpledialog.askstring("Campo", "Nome do campo", parent=self)
        if not name:
            return
        section = self.section.get()
        target = (
            self.draft[section]
            if section == "proposal"
            else self.draft[section]["columns"]
        )
        if name in target:
            self.error.set("Já existe um campo com esse nome.")
            return
        target[name] = (
            {
                "kind": "manual",
                "type": "text",
                "value": "",
                "source": "main",
                "sheet": "",
                "address": "",
            }
            if section == "proposal"
            else name
        )
        self.refresh_fields()

    def remove_field(self):
        section, key = self.section.get(), self.field_name.get()
        target = (
            self.draft[section]
            if section == "proposal"
            else self.draft[section]["columns"]
        )
        target.pop(key, None)
        self.refresh_fields()

    def open_book(self):
        path = filedialog.askopenfilename(
            parent=self, filetypes=[("Excel", "*.xlsx *.xlsm *.xls")]
        )
        if not path:
            return
        try:
            if self.book:
                self.book.close()
            self.book = Workbook(path)
            self.sheet["values"] = self.book.sheets
            self.sheet.current(0)
            self.preview()
        except Exception as exc:
            self.error.set(str(exc))

    def preview(self):
        from openpyxl.utils import get_column_letter

        self.grid.delete(*self.grid.get_children())
        sheet = self.sheet.get()
        rows, cols = self.book.size(sheet)
        self.grid["columns"] = ["row"] + [
            get_column_letter(c) for c in range(1, min(cols, 30) + 1)
        ]
        for c in self.grid["columns"]:
            self.grid.heading(c, text=c)
            self.grid.column(c, width=85, minwidth=50, stretch=False)
        for row in range(1, min(rows, 100) + 1):
            values = []
            for col in range(1, min(cols, 30) + 1):
                try:
                    values.append(
                        str(self.book.get(sheet, row, col, self.draft["rules"]))
                    )
                except ValueError as exc:
                    values.append(str(exc))
            self.grid.insert("", "end", iid=str(row), values=[row] + values)

    def pick_cell(self, event):
        row = self.grid.identify_row(event.y)
        column = self.grid.identify_column(event.x)
        if not row or not column or int(column[1:]) < 2:
            return
        from openpyxl.utils import get_column_letter

        col = int(column[1:]) - 1
        self.map_vars["sheet"].set(self.sheet.get())
        if self.section.get() == "proposal":
            self.map_vars["address"].set(f"{get_column_letter(col)}{row}")
            self.map_vars["kind"].set("cell")
        else:
            self.map_vars["column"].set("@" + get_column_letter(col))

    def test_mapping(self):
        mapper = None
        try:
            self.apply_mapping()
            self.collect()
            mapper = Mapper(self.draft, self.book.path if self.book else "")
            result = (
                mapper.proposal()
                if self.section.get() == "proposal"
                else mapper.table(self.draft[self.section.get()])
            )
            dialog = tk.Toplevel(self)
            dialog.title("Resultado do mapeamento")
            text = tk.Text(dialog, width=100, height=25)
            text.pack(fill="both", expand=True)
            text.insert(
                "1.0", json.dumps(result, ensure_ascii=False, indent=2, default=str)
            )
            text.configure(state="disabled")
        except Exception as exc:
            self.error.set(str(exc))
        finally:
            if mapper:
                mapper.close()

    def test_connection(self, protocol):
        try:
            from .mail import connection

            self.collect()
            password = getattr(
                self, protocol + "_password"
            ).get() or self.store.password(self.draft["id"], protocol)
            client = connection(
                self.draft["tools"]["email"][protocol], password, protocol
            )
            client.quit() if protocol == "smtp" else client.logout()
            self.error.set("Ligação autenticada com sucesso. Nenhuma mensagem enviada.")
        except Exception as exc:
            self.error.set(
                f"Não foi possível ligar: {type(exc).__name__}. Verifique servidor e credenciais."
            )
