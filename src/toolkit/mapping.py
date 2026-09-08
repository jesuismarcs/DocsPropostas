"""Read-only workbook mapping, with explicit conversion and cell provenance."""

from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter


class Workbook:
    def __init__(self, path):
        self.path = str(path)
        self.legacy = Path(path).suffix.lower() == ".xls"
        if self.legacy:
            import xlrd

            self.values = xlrd.open_workbook(str(path), formatting_info=True)
            self.formulas = None
            self.sheets = self.values.sheet_names()
        else:
            self.values = openpyxl.load_workbook(path, data_only=True)
            self.formulas = openpyxl.load_workbook(path, data_only=False)
            self.sheets = self.values.sheetnames

    def close(self):
        if not self.legacy:
            self.values.close()
            self.formulas.close()
        else:
            self.values.release_resources()

    def size(self, sheet):
        if sheet not in self.sheets:
            raise ValueError(f"Folha não encontrada: {sheet}")
        ws = self.values.sheet_by_name(sheet) if self.legacy else self.values[sheet]
        return (ws.nrows, ws.ncols) if self.legacy else (ws.max_row, ws.max_column)

    def hidden(self, sheet, row):
        if self.legacy:
            info = self.values.sheet_by_name(sheet).rowinfo_map.get(row - 1)
            return bool(info and info.hidden)
        return self.values[sheet].row_dimensions[row].hidden

    def get(self, sheet, row, col, rules):
        label = f"{sheet}!{get_column_letter(col)}{row}"
        self.size(sheet)
        if self.legacy:
            import xlrd

            ws = self.values.sheet_by_name(sheet)
            for r1, r2, c1, c2 in ws.merged_cells:
                if r1 <= row - 1 < r2 and c1 <= col - 1 < c2:
                    if rules.get("merged") == "fill":
                        row, col = r1 + 1, c1 + 1
                    elif rules.get("merged") == "error":
                        raise ValueError(f"{label}: célula unida")
            if row > ws.nrows or col > ws.ncols:
                return ""
            cell = ws.cell(row - 1, col - 1)
            if cell.ctype == xlrd.XL_CELL_ERROR:
                raise ValueError(f"{label}: erro Excel")
            if cell.ctype == xlrd.XL_CELL_DATE:
                return xlrd.xldate_as_datetime(cell.value, self.values.datemode)
            if cell.ctype == xlrd.XL_CELL_NUMBER:
                xf = self.values.xf_list[ws.cell_xf_index(row - 1, col - 1)]
                fmt = self.values.format_map[xf.format_key].format_str
                if re.fullmatch(r"0+", fmt):
                    return str(int(cell.value)).zfill(len(fmt))
            return cell.value
        ws = self.values[sheet]
        for area in ws.merged_cells.ranges:
            if ws.cell(row, col).coordinate in area:
                if rules.get("merged") == "fill":
                    row, col = area.min_row, area.min_col
                elif rules.get("merged") == "error":
                    raise ValueError(f"{label}: célula unida")
        c = ws.cell(row, col)
        f = self.formulas[sheet].cell(row, col)
        if f.data_type == "f" and c.value is None:
            raise ValueError(f"{label}: recalcule e guarde o livro no Excel.")
        if c.data_type == "e":
            raise ValueError(f"{label}: erro Excel {c.value}")
        value = c.value
        if isinstance(value, (int, float)) and re.fullmatch(r"0+", c.number_format):
            return str(int(value)).zfill(len(c.number_format))
        return "" if value is None else value

    def resolve(self, mapping):
        sheet, address = mapping.get("sheet", ""), mapping.get("address", "")
        kind = mapping.get("kind", "range")
        if kind == "name":
            if self.legacy:
                names = self.values.name_map.get(address.lower(), [])
                if len(names) != 1:
                    raise ValueError(f"Nome definido ausente ou ambíguo: {address}")
                ws, r1, r2, c1, c2 = names[0].area2d(clipped=False)
                return (
                    ws.name,
                    f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}",
                )
            name = self.values.defined_names.get(address)
            if name is None:
                raise ValueError(f"Nome definido inexistente: {address}")
            destinations = list(name.destinations)
            if len(destinations) != 1:
                raise ValueError("Escolha um nome que identifique um único intervalo.")
            sheet, address = destinations[0]
        elif kind == "table":
            if self.legacy:
                raise ValueError("Tabelas estruturadas exigem .xlsx ou .xlsm.")
            if sheet not in self.sheets or address not in self.values[sheet].tables:
                raise ValueError(f"Tabela não encontrada: {sheet}!{address}")
            address = self.values[sheet].tables[address].ref
        self.size(sheet)
        return sheet, address


def convert(value, kind, rules):
    if value == "" or value is None:
        return ""
    if kind == "text":
        return str(value).strip()
    if kind == "number":
        if isinstance(value, (int, float, Decimal)):
            return float(value)
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        thousands = rules.get("thousands", ".")
        if thousands:
            text = text.replace(thousands, "")
        try:
            return float(Decimal(text.replace(rules.get("decimal", ","), ".")))
        except InvalidOperation as exc:
            raise ValueError(f"Número inválido: {value}") from exc
    if kind == "date":
        if isinstance(value, (date, datetime)):
            return value.strftime("%Y-%m-%d")
        return datetime.strptime(str(value).strip(), rules["date_format"]).strftime(
            "%Y-%m-%d"
        )
    return value


class Mapper:
    def __init__(self, profile, main_path):
        self.profile, self.main_path, self.books = profile, main_path, {}

    def book(self, mapping):
        source = mapping.get("source", "main")
        path = (
            self.main_path
            if source == "main" and self.main_path
            else self.profile["sources"].get(source)
        ) or self.main_path
        if not path:
            raise ValueError(f"Selecione o ficheiro de origem: {source}")
        if path not in self.books:
            self.books[path] = Workbook(path)
        return self.books[path]

    def close(self):
        for book in self.books.values():
            book.close()

    def field(self, mapping):
        kind, dtype = mapping["kind"], mapping.get("type", "text")
        if kind in ("constant", "manual"):
            value = mapping.get("value", "")
            if dtype == "list" and isinstance(value, str):
                value = [x.strip() for x in value.splitlines() if x.strip()]
            elif dtype == "table" and isinstance(value, str):
                import csv
                import io

                value = (
                    list(csv.DictReader(io.StringIO(value), delimiter="\t"))
                    if value.strip()
                    else []
                )
        else:
            book = self.book(mapping)
            sheet, address = book.resolve(mapping)
            c1, r1, c2, r2 = range_boundaries(address)
            rows = [
                [
                    book.get(sheet, r, c, self.profile["rules"])
                    for c in range(c1, c2 + 1)
                ]
                for r in range(r1, r2 + 1)
                if self.profile["rules"].get("hidden") != "skip"
                or not book.hidden(sheet, r)
            ]
            if dtype == "table":
                if not rows:
                    return []
                headers = [str(x) for x in rows[0]]
                if len(set(headers)) != len(headers) or "" in headers:
                    raise ValueError("Cabeçalhos da tabela vazios ou duplicados.")
                value = [
                    dict(zip(headers, row))
                    for row in rows[1:]
                    if any(x != "" for x in row)
                ]
            elif dtype == "list":
                value = [str(x) for row in rows for x in row if x != ""]
            else:
                if len(rows) != 1 or len(rows[0]) != 1:
                    raise ValueError("Um campo simples exige uma única célula.")
                value = rows[0][0]
        if mapping.get("required") and (value == "" or value == []):
            raise ValueError("Campo obrigatório vazio.")
        return convert(value, dtype, self.profile["rules"])

    def proposal(self):
        data, errors = {}, []
        for key, mapping in self.profile["proposal"].items():
            try:
                data[key] = self.field(mapping)
            except Exception as exc:
                errors.append(
                    f"{key} ({mapping.get('sheet', '')}!{mapping.get('address', '')}): {exc}"
                )
        if errors:
            raise ValueError("\n".join(errors))
        return data

    def table(self, mapping):
        book = self.book(mapping)
        sheet, address = book.resolve(mapping)
        maxrow, maxcol = book.size(sheet)
        h, first, last, startcol = (
            int(mapping.get("header_row", 1)),
            int(mapping.get("first_row", 2)),
            int(mapping.get("last_row", 0)) or maxrow,
            1,
        )
        if address:
            startcol, h, maxcol, last = range_boundaries(address)
            first = h + 1
        headers = [
            str(book.get(sheet, h, c, self.profile["rules"])).strip()
            for c in range(startcol, maxcol + 1)
        ]
        indices = {}
        for key, label in mapping["columns"].items():
            if label.startswith("@"):  # Explicit physical column, independent of label.
                from openpyxl.utils import column_index_from_string

                indices[key] = column_index_from_string(label[1:])
            elif headers.count(label) == 1:
                indices[key] = startcol + headers.index(label)
            else:
                raise ValueError(
                    f"{sheet}: cabeçalho '{label}' ausente ou duplicado ({key})."
                )
        result = []
        for row in range(first, last + 1):
            if self.profile["rules"].get("hidden") == "skip" and book.hidden(
                sheet, row
            ):
                continue
            record = {
                k: book.get(sheet, row, col, self.profile["rules"])
                for k, col in indices.items()
            }
            if not any(v != "" for v in record.values()):
                if self.profile["rules"].get("blank_rows") == "stop":
                    break
                continue
            for k, v in record.items():
                try:
                    record[k] = convert(
                        v,
                        mapping.get("types", {}).get(k, "text"),
                        self.profile["rules"],
                    )
                except Exception as exc:
                    raise ValueError(
                        f"{sheet}!{get_column_letter(indices[k])}{row}: {exc}"
                    ) from exc
            result.append(record)
        return result
